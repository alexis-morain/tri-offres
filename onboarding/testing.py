"""Test fixtures of the onboarding: the summary screen's test case and URLs, the block pages' form data, and for the
CV screen, PDF and DOCX files built in memory and archives with doctored headers."""
import io
import struct
import zipfile
import zlib

import docx
from pypdf import PdfReader, PdfWriter

from accounts.models import OnboardingEvent, Profile, SummaryStatus
from llm import fake
from onboarding import summary
from onboarding.tests import CULTURE_FULL, SUMMARY, QuestionnaireTestCase

MY_SEARCH = "/ma-recherche/"
WRITE = "/synthese/rediger/"
CORRECT = "/synthese/corriger/"
TEXT = "/synthese/texte/"
VALIDATE = "/synthese/valider/"

GOOD = {
    "sentences": [
        {"text": "J’ai compris que vous cherchez un poste où vous parlez aux publics.", "sources": ["seek_text"]},
        {"text": "Vous visez la médiation et la communication, à Lyon ou autour.", "sources": ["postes", "location"]},
        {"text": "Le sens du travail et l’équipe comptent le plus pour vous.", "sources": ["values"]},
        {"text": "Vous aimeriez travailler en maison d’édition ou en bibliothèque.", "sources": ["structures"]},
    ],
    "questions": [{"text": "Je n’ai pas bien saisi si l’accueil du public doit être au cœur du poste.",
                   "sources": ["seek_text"]}],
}


class SummaryTestCase(QuestionnaireTestCase):
    def setUp(self):
        super().setUp()
        fake.reset()
        self.addCleanup(fake.reset)
        fake.set_response(summary.TASK, GOOD)
        self.walk(CULTURE_FULL)

    def update(self, **fields):
        Profile.objects.filter(user=self.user).update(**fields)

    def page(self, path=SUMMARY):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        return response.content.decode()

    def ready(self, **fields):
        """A summary written by the fake engine, then `fields` set on the profile."""
        self.client.post(WRITE)
        self.assertEqual(self.profile.summary_status, SummaryStatus.READY)
        if fields:
            self.update(**fields)

    def card(self, html, tone):
        start = html.index(f'class="card card--{tone}')
        return html[start:html.index("</article>", start)]

    def event_names(self):
        return list(OnboardingEvent.objects.filter(user=self.user).values_list("name", flat=True))


def block_url(slug):
    return f"{MY_SEARCH}{slug}/"


def prefixed(slug, data):
    return {f"{slug}-{key}": value for key, value in data.items()}

# A CV long enough not to look scanned, with contact details the redaction must remove.
CV_LINES = [
    "Nora Testard",
    "12 rue des Essais",
    "99000 Exempleville",
    "nora.testard@example.fr, 01 99 00 00 00",
    "Médiatrice culturelle, Musée de l’Exemple, 2019-2023",
    "Chargée d’accueil des publics, Bibliothèque municipale d’Exempleville, 2016-2019",
    "Compétences : médiation, écriture de dossiers, accueil de groupes scolaires, gestion de projet",
    "Master 2 Médiation culturelle, Université d’Exemple, 2016",
]


def _pdf_string(text):
    raw = text.encode("cp1252", errors="replace")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def make_pdf(pages=None):
    """A PDF with one page per item of `pages` (a list of lines; an empty list makes a page without text)."""
    pages = [CV_LINES] if pages is None else pages
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
               3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"}
    kids = []
    for index, lines in enumerate(pages):
        page_id, content_id = 4 + 2 * index, 5 + 2 * index
        kids.append(f"{page_id} 0 R".encode())
        stream = b"BT /F1 11 Tf 50 800 Td 14 TL " + b" ".join(_pdf_string(line) + b" Tj T*" for line in lines) + b" ET"
        objects[page_id] = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R"
                            b" >> >> /Contents " + str(content_id).encode() + b" 0 R >>")
        objects[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    objects[2] = b"<< /Type /Pages /Kids [" + b" ".join(kids) + b"] /Count %d >>" % len(pages)
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = out.tell()
        out.write(b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n")
    xref = out.tell()
    size = max(objects) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    for number in range(1, size):
        out.write(b"%010d 00000 n \n" % offsets[number])
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref))
    return out.getvalue()


def make_encrypted_pdf(user_password="secret"):
    """An encrypted PDF. With an empty user password it opens without one: only its usage is restricted."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf())))
    writer.encrypt(user_password=user_password, owner_password="owner", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_pdf_bomb(megabytes=1024):
    """A one-page PDF whose content stream inflates to `megabytes` of zeros (about 1 MB on disk per GB)."""
    compressor = zlib.compressobj(9)
    chunk = b"\0" * (1024 * 1024)
    stream = b"".join(compressor.compress(chunk) for _ in range(megabytes)) + compressor.flush()
    head = make_pdf([["x"]])
    marker = head.index(b"5 0 obj")
    body = (b"5 0 obj\n<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream) + stream
            + b"\nendstream\nendobj\n")
    pdf = head[:marker] + body
    # No valid xref: pypdf rebuilds it by scanning, as it does for many real-world files.
    return pdf + b"trailer\n<< /Size 6 /Root 1 0 R >>\n%%EOF\n"


def make_docx(paragraphs=None, table=None):
    """A .docx with these paragraphs, and an optional table (a list of rows of cell texts)."""
    document = docx.Document()
    for text in CV_LINES if paragraphs is None else paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in zip(grid.rows, table):
            for cell, text in zip(row.cells, cells):
                cell.text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def rezip(data, extra=(), drop=(), stream_entry=None):
    """A copy of the archive `data`, without the `drop` names, with (name, bytes) entries of `extra` added, and
    with `stream_entry` = (name, megabytes) written as that many megabytes of zeros, in 1 MB chunks."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename not in drop:
                target.writestr(info.filename, source.read(info.filename))
        for name, content in extra:
            target.writestr(name, content)
        if stream_entry:
            name, megabytes = stream_entry
            with target.open(name, "w", force_zip64=True) as handle:
                for _ in range(megabytes):
                    handle.write(b"\0" * (1024 * 1024))
    return out.getvalue()


def patch_central_directory(data, name, *, flag_bits=None, compress_size=None, file_size=None):
    """The archive `data` with the central directory record of `name` doctored (sizes that lie, an encryption flag).
    The local header gets the same sizes, so the archive still opens."""
    data = bytearray(data)
    position = data.find(b"PK\x01\x02")
    while position != -1:
        name_length, extra_length, comment_length = struct.unpack_from("<HHH", data, position + 28)
        if data[position + 46:position + 46 + name_length].decode() == name:
            local = struct.unpack_from("<I", data, position + 42)[0]
            if flag_bits is not None:
                struct.pack_into("<H", data, position + 8, flag_bits)
                struct.pack_into("<H", data, local + 6, flag_bits)
            if compress_size is not None:
                struct.pack_into("<I", data, position + 20, compress_size)
                struct.pack_into("<I", data, local + 18, compress_size)
            if file_size is not None:
                struct.pack_into("<I", data, position + 24, file_size)
                struct.pack_into("<I", data, local + 22, file_size)
            return bytes(data)
        position = data.find(b"PK\x01\x02", position + 46 + name_length + extra_length + comment_length)
    raise KeyError(name)
